"""Convert the filled abstract Markdown (a deliberately small Markdown subset) into a .docx that
mirrors the official Round 1 Google-Doc template, so the team can paste it into their copy."""
from __future__ import annotations

import pathlib
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text):
    text = text.replace("\\|", "|")
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(tok)


def build(md_path: pathlib.Path, out_path: pathlib.Path, fig_dir: pathlib.Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    usable = sec.page_width - sec.left_margin - sec.right_margin
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        mt = re.match(r'<p class="(t\d)">(.*)</p>', line.strip())
        if mt:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(mt.group(2))
            r.bold = True
            r.font.size = Pt({"t1": 16, "t2": 15, "t3": 13}.get(mt.group(1), 13))
            i += 1
            continue
        mb = re.match(r"^( *)- (.*)", line)
        if mb:
            level = 1 if len(mb.group(1)) < 4 else 2
            item = [mb.group(2)]
            i += 1
            while (i < len(lines) and lines[i].strip() and lines[i].startswith(" ")
                   and not re.match(r"^ *- ", lines[i])):
                item.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
            add_runs(p, " ".join(item))
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            if m:
                img = fig_dir / pathlib.Path(m.group(2)).name
                if img.exists():
                    doc.add_picture(str(img), width=usable)
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, txt in enumerate(cells):
                    cell = t.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], txt)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        if r == 0:
                            run.bold = True
            continue
        elif line.startswith("- "):
            item = [line[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                item.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, " ".join(item))
            continue
        else:
            para = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"(#|!\[|\||- )", lines[i]):
                para.append(lines[i].strip())
                i += 1
            txt = " ".join(para)
            p = doc.add_paragraph()
            if txt.startswith("*Figure"):
                add_runs(p, txt)
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
            else:
                add_runs(p, txt)
            continue
        i += 1
    doc.save(str(out_path))
